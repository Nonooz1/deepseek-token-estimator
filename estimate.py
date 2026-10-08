#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Token 用量估算器（多格式文件版）
—— 上传文件到对话前，先看看它要吃掉多少 token

支持格式：
  压缩包  ：.zip .7z .rar(需外部工具) .tar .tar.gz/.tgz .tar.bz2 .tar.xz
            支持嵌套压缩包（默认最多 3 层）
  Keil    ：.uvprojx .uvoptx .uvproj .uvopt .sct .a51 等 MDK / C51 工程文件
            自动跳过编译产物（.o/.obj/.axf/.hex/.map/.lst/Objects//Output//LST//si/ 等）
  代码    ：.c .h .cpp .hpp .s .asm .a51 .inc .ld .icf 等
  纯文本  ：.txt .md .py .js .ts .json .csv .log .yaml .xml .html ...
  Word    ：.docx / .docm        （需要 python-docx）
  PDF     ：.pdf                 （需要 pypdf）
  Excel   ：.xlsx / .xlsm        （需要 openpyxl）

用法：
  python estimate.py 工程.zip
  python estimate.py 工程.zip --all-files      # 列出包内每个文件
  python estimate.py 工程.zip --all            # 连编译产物一起算
  python estimate.py 报告.docx 数据.xlsx
  python estimate.py C:\\某个目录
  python estimate.py 工程.zip -q               # 只输出总 token 数
"""
import argparse
import io
import os
import re
import shutil
import subprocess
import sys
import tarfile
import tempfile
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))

# 由 --force-text 打开：忽略二进制嗅探，一律按文本读
FORCE_TEXT = False

# 由 --all 打开：连编译产物（.o/.axf/.map/Objects 目录等）一起算
INCLUDE_ARTIFACTS = False

# 由 --exclude 追加：路径中任意一段等于这些名字就跳过（大小写不敏感）
EXCLUDE_PATS = []


# ---------------------------------------------------------------- 终端宽度处理
def disp_width(s: str) -> int:
    """中日韩字符在终端占 2 列"""
    w = 0
    for ch in s:
        o = ord(ch)
        if (0x1100 <= o <= 0x115F or 0x2E80 <= o <= 0xA4CF or
                0xAC00 <= o <= 0xD7A3 or 0xF900 <= o <= 0xFAFF or
                0xFE30 <= o <= 0xFE6F or 0xFF00 <= o <= 0xFF60 or
                0xFFE0 <= o <= 0xFFE6 or 0x20000 <= o <= 0x3FFFD):
            w += 2
        else:
            w += 1
    return w


def pad(s: str, width: int, align: str = "left") -> str:
    gap = max(0, width - disp_width(s))
    return s + " " * gap if align == "left" else " " * gap + s


def trunc(s: str, width: int) -> str:
    """按显示宽度截断，超长加省略号"""
    if disp_width(s) <= width:
        return s
    out, w = "", 0
    for ch in s:
        cw = disp_width(ch)
        if w + cw > width - 1:
            break
        out += ch
        w += cw
    return out + "…"


def human(n):
    return f"{n:,}" if isinstance(n, int) and n >= 1000 else str(n)


def table(headers, widths, aligns, rows, indent=0):
    """打印一张对齐的表（中英文混排按显示宽度对齐）"""
    pre = " " * indent
    print(pre + "  ".join(pad(h, w, a) for h, w, a in zip(headers, widths, aligns)))
    print(pre + "-" * (sum(widths) + 2 * (len(widths) - 1)))
    for r in rows:
        print(pre + "  ".join(pad(trunc(str(c), w), w, a)
                              for c, w, a in zip(r, widths, aligns)))


# ---------------------------------------------------------------- 分词器加载
def load_tokenizer():
    """返回 (tokenizer, 后端名)"""
    errors = []
    try:
        import transformers
        tok = transformers.PreTrainedTokenizerFast.from_pretrained(HERE)
        return tok, f"transformers {transformers.__version__} (fast)"
    except Exception as e:
        errors.append(f"transformers: {e}")
    try:
        from tokenizers import Tokenizer
        return Tokenizer.from_file(os.path.join(HERE, "tokenizer.json")), "tokenizers (lightweight)"
    except Exception as e:
        errors.append(f"tokenizers: {e}")
    sys.stderr.write("无法加载分词器，请先：pip install transformers\n" + "\n".join(errors) + "\n")
    sys.exit(1)


def count(tok, text: str) -> int:
    if hasattr(tok, "convert_ids_to_tokens"):
        return len(tok.encode(text, add_special_tokens=False))
    return len(tok.encode(text).ids)


# ---------------------------------------------------------------- 扩展名分类
# Keil MDK / 嵌入式工程相关
KEIL_EXT = {
    ".uvprojx", ".uvoptx", ".uvproj", ".uvopt",      # uVision 5 / 4 工程与选项
    ".sct", ".scf", ".icf", ".ld", ".lds",           # 分散加载 / 链接脚本
    ".gpdsc", ".pdsc", ".scvd",                      # Pack / 组件描述
    ".inc", ".a51", ".c51", ".src",                  # 汇编与包含文件
    ".dbgconf", ".jflash", ".svd",                   # 调试配置
}

CODE_EXT = {
    ".c", ".h", ".cpp", ".cc", ".cxx", ".hpp", ".hh", ".hxx", ".inl",
    ".s", ".asm", ".sx", ".as",
    ".cs", ".java", ".kt", ".go", ".rs", ".rb", ".php", ".swift", ".m", ".mm",
    ".py", ".pyw", ".js", ".mjs", ".cjs", ".ts", ".tsx", ".jsx", ".vue",
    ".sql", ".sh", ".bash", ".bat", ".cmd", ".ps1", ".vbs", ".tcl", ".cmake",
    ".makefile", ".mk", ".gradle", ".vb", ".vhd", ".v", ".sv",
}

TEXT_EXT = {
    ".txt", ".md", ".markdown", ".rst", ".log", ".csv", ".tsv",
    ".json", ".jsonl", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf", ".env",
    ".html", ".htm", ".xml", ".css", ".scss", ".less", ".properties",
    ".map", ".lst", ".srec", ".hex", ".ihex",
    ".editorconfig", ".gitignore", ".gitattributes",
} | CODE_EXT | KEIL_EXT

# 二进制扩展名
BINARY_EXT = {
    ".exe", ".dll", ".so", ".dylib", ".pyd", ".pyc", ".pyo", ".class",
    ".rar", ".7z", ".cab", ".iso", ".msi", ".lzh", ".arj",
    ".png", ".jpg", ".jpeg", ".gif", ".bmp", ".ico", ".webp", ".tif", ".tiff", ".psd",
    ".mp3", ".mp4", ".avi", ".mkv", ".mov", ".wav", ".flac", ".aac", ".wmv", ".webm",
    ".ttf", ".otf", ".woff", ".woff2", ".eot",
    ".db", ".sqlite", ".sqlite3", ".mdb", ".accdb",
    ".bin", ".dat", ".img", ".raw", ".elf", ".flm", ".jar", ".war", ".apk", ".ipa",
    ".nupkg", ".whl", ".egg", ".o", ".obj", ".a", ".axf", ".lib", ".crf",
}

# 压缩包
TAR_EXT = {".tar", ".tgz", ".tbz2", ".txz",
           ".tar.gz", ".tar.bz2", ".tar.xz", ".tar.zst"}
ARCHIVE_EXT = {".zip", ".7z", ".rar"} | TAR_EXT

# .rar 需要外部工具（unrar / 7-Zip / Bandizip）。纯 Python 做不到：
# Windows 自带的 bsdtar 只能「列出」rar 成员，解出来的数据是错的
# （实测要 5192 字节只给 53 字节），所以不能拿它当后端。
RAR_TOOLS = [
    ("unrar", ["{exe}", "x", "-y", "-idq", "{src}", "{dest}" + os.sep]),
    ("unrar-free", ["{exe}", "x", "-y", "{src}", "{dest}"]),
    ("7z", ["{exe}", "x", "-y", "-bd", "-o{dest}", "{src}"]),
    ("7za", ["{exe}", "x", "-y", "-bd", "-o{dest}", "{src}"]),
    ("7zz", ["{exe}", "x", "-y", "-bd", "-o{dest}", "{src}"]),
    ("bz", ["{exe}", "x", "-y", "-o:{dest}", "{src}"]),
    ("Rar.exe", ["{exe}", "x", "-y", "-ibck", "{src}", "{dest}" + os.sep]),
]
RAR_TOOL_DIRS = [
    r"C:\Program Files\7-Zip", r"C:\Program Files (x86)\7-Zip",
    r"C:\Program Files\WinRAR", r"C:\Program Files (x86)\WinRAR",
    r"C:\Program Files\Bandizip", r"C:\Program Files (x86)\Bandizip",
    r"C:\Program Files\NanaZip", r"C:\Program Files (x86)\NanaZip",
]
# 超过这个大小就不自动解压 rar，避免临时目录被撑爆
RAR_EXTRACT_LIMIT = 200 * 1024 * 1024

# 编译/链接产物：是文本但没营养，默认跳过（加 --all 才包含）
# 参考一个真实 Keil C51 工程（Nuvoton MS51）里出现的东西
ARTIFACT_EXT = {
    # 目标文件与镜像
    ".o", ".obj", ".axf", ".elf", ".hex", ".ihex", ".bin", ".lib", ".a",
    # Keil 编译器/链接器中间产物
    ".crf", ".d", ".dep", ".lnp", ".iex", ".lst", ".map", ".sbr", ".__i",
    # 界面布局与备份（uVision 5 是 .uvguix.<用户>，4 是 .uvgui.<用户>）
    ".uvgui", ".uvguix", ".bak", ".orig", ".rej", ".tmp",
    # 构建日志（.build_log.htm）
    ".htm", ".html",
    # Source Insight 的索引（配合 Keil 用得很普遍）
    ".iab", ".iad", ".imb", ".imd", ".pfi", ".po", ".pr", ".pri", ".ps",
    ".searchresults", ".wk3",
}

# 产物文件名特征（Keil / Source Insight 常见的生成物）
# 注意：uVision 5 写 *.uvguix.<用户名>，uVision 4 写 *.uvgui.<用户名>，
# 真实扩展名是「用户名」，所以必须按文件名匹配，用 splitext 会漏。
ARTIFACT_NAME_RE = re.compile(
    r"\.uvgui[x]?\.|\.build_log\.|\.dep$|\.crf$|\.lnp$|\.iex$|"
    r"\.__i$|\.sbr$|\.bak\d*$|\.old$|\.orig$", re.I)

# 这些目录里基本都是产物
SKIP_DIRS = {
    # Keil uVision
    "objects", "listings", "output", "lst", "debugconfig",
    # Source Insight 索引目录
    "si",
    # 通用
    "__pycache__", ".git", ".svn", ".vscode", ".vs", ".idea",
}

# 目录跳过时的豁免名单：万一用户自己有个 output/ 放源码，别整目录丢掉。
# 只豁免明确的源码/工程文件扩展名，产物扩展名（.lst/.obj/.hex…）不在其中。
SOURCE_EXT = {
    ".c", ".h", ".cpp", ".cxx", ".cc", ".hpp", ".hh", ".hxx", ".inl",
    ".a51", ".a51c", ".s", ".asm", ".src", ".inc", ".sct", ".ld", ".icf",
    ".uvproj", ".uvprojx", ".uvopt", ".uvoptx",
    ".ini", ".txt", ".md", ".cfg",
}


class BinaryFile(Exception):
    """内容看起来是二进制，不是文本"""


def is_archive(name):
    return archive_kind(name) is not None


def archive_kind(name):
    """返回 'zip' / '7z' / 'rar' / 'tar' / None"""
    low = name.lower()
    if low.endswith(".zip"):
        return "zip"
    if low.endswith(".7z"):
        return "7z"
    if low.endswith(".rar"):
        return "rar"
    if any(low.endswith(e) for e in TAR_EXT):
        return "tar"
    return None


def find_rar_tool():
    """找一个能解 rar 的外部程序。找不到返回 None。"""
    for exe, argv in RAR_TOOLS:
        found = shutil.which(exe)
        if not found:
            for d in RAR_TOOL_DIRS:
                cand = os.path.join(d, exe)
                if os.path.isfile(cand):
                    found = cand
                    break
        if found:
            return found, argv
    return None


def extract_rar_to_temp(path):
    """把 rar 解到临时目录，返回 (临时目录对象, 错误信息)。
    调用方负责在 with 块里用完就清理。"""
    tool = find_rar_tool()
    if tool is None:
        return None, ("读 .rar 需要外部程序（unrar 或 7-Zip）。"
                      "装一个 7-Zip 后会自动识别；也可以先手动解压成 zip。")
    exe, argv = tool
    dest = tempfile.mkdtemp(prefix="tokrar_")
    cmd = [a.format(exe=exe, src=path, dest=dest) for a in argv]
    try:
        r = subprocess.run(cmd, capture_output=True, timeout=300)
    except subprocess.TimeoutExpired:
        shutil.rmtree(dest, ignore_errors=True)
        return None, "解压超时（超过 5 分钟）"
    except Exception as e:
        shutil.rmtree(dest, ignore_errors=True)
        return None, f"解压失败：{e}"
    if r.returncode != 0:
        shutil.rmtree(dest, ignore_errors=True)
        msg = (r.stderr or r.stdout or b"").decode("utf-8", "replace").strip()
        return None, f"解压失败（{os.path.basename(exe)} 退出码 {r.returncode}）{msg[:200]}"
    return dest, ""


def base_name(virtual):
    """从虚拟路径里取文件名（兼容 Windows 反斜杠与嵌套压缩包的 ! 分隔）"""
    return virtual.replace("!", "/").replace("\\", "/").rsplit("/", 1)[-1]


def in_skip_dir(virtual):
    norm = virtual.replace("!", "/").replace("\\", "/")
    parts = [p.lower() for p in norm.split("/")[:-1]]
    if not any(p in SKIP_DIRS for p in parts):
        return False
    # 目录撞名保护：明确是源码/工程文件的不按目录跳过
    ext = os.path.splitext(base_name(virtual))[1].lower()
    return ext not in SOURCE_EXT


def is_artifact(virtual):
    """判断是不是编译产物"""
    name = base_name(virtual)
    if ARTIFACT_NAME_RE.search(name):
        return True
    ext = os.path.splitext(name)[1].lower()
    return ext in ARTIFACT_EXT


def in_excluded(virtual):
    """--exclude 指定的目录名（整段匹配，大小写不敏感）"""
    if not EXCLUDE_PATS:
        return False
    norm = virtual.replace("!", "/").replace("\\", "/")
    parts = [p.lower() for p in norm.split("/")[:-1]]
    return any(p in EXCLUDE_PATS for p in parts)


# ---------------------------------------------------------------- 编码嗅探
def _sniff_bom(head):
    """按 BOM 判断编码。注意 UTF-32 的 BOM 以 UTF-16 的 BOM 开头，必须先判 UTF-32。"""
    for bom, enc in (
        (b"\xff\xfe\x00\x00", "utf-32"),
        (b"\x00\x00\xfe\xff", "utf-32"),
        (b"\xef\xbb\xbf", "utf-8-sig"),
        (b"\xff\xfe", "utf-16"),
        (b"\xfe\xff", "utf-16"),
    ):
        if head.startswith(bom):
            return enc
    return None


def _sniff_utf16_no_bom(head):
    """无 BOM 的 UTF-16 猜测：ASCII 为主的文本，NUL 会规律地出现在固定奇偶位。"""
    if len(head) < 8:
        return None
    half = len(head) // 2
    even_nul = sum(1 for i in range(0, 2 * half, 2) if head[i] == 0)
    odd_nul = sum(1 for i in range(1, 2 * half, 2) if head[i] == 0)
    if odd_nul / half > 0.7:
        return "utf-16-le"
    if even_nul / half > 0.7:
        return "utf-16-be"
    return None


def looks_binary_bytes(name, data, probe=8192):
    """判断字节内容是不是二进制。
    注意：latin-1 解码永远不会抛异常，所以「解码成功」不能证明是文本，
    必须另外闻一下字节。否则随机二进制会被算成几十万 token 的噪声。"""
    ext = os.path.splitext(name)[1].lower()
    if ext in BINARY_EXT:
        return True, f"{ext} 是二进制格式"

    head = data[:probe]
    if not head:
        return False, ""

    # 带 BOM 的 UTF-16/32 天然含 NUL，先放行
    if _sniff_bom(head) or _sniff_utf16_no_bom(head):
        return False, ""

    if b"\x00" in head:
        return True, "含 NUL 字节"

    ctrl = sum(1 for b in head if b < 9 or 13 < b < 32 or b == 127)
    if ctrl / len(head) > 0.30:
        return True, f"控制字符占比 {ctrl / len(head):.0%}"

    return False, ""


def decode_bytes(name, data):
    """解码字节。顺序：BOM → 无 BOM UTF-16 猜测 → utf-8 → gbk → gb18030 → latin-1。
    解码前先做二进制嗅探，避免把二进制当文本。"""
    binary, why = looks_binary_bytes(name, data)
    if binary:
        raise BinaryFile(why)

    head = data[:8192]
    enc = _sniff_bom(head) or _sniff_utf16_no_bom(head)
    if enc:
        try:
            return data.decode(enc), enc
        except (UnicodeDecodeError, UnicodeError):
            pass

    for enc in ("utf-8", "utf-8-sig", "gbk", "gb18030", "latin-1"):
        try:
            return data.decode(enc), enc
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise ValueError("无法识别文件编码")


def read_text_file(path):
    with open(path, "rb") as fh:
        data = fh.read()
    return decode_bytes(os.path.basename(path), data)


# ---------------------------------------------------------------- 结构化抽取
def extract_docx(path):
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for tbl in doc.tables:
        for row in tbl.rows:
            parts.append("\t".join(c.text for c in row.cells))
    return "\n".join(parts)


def extract_pdf(path):
    from pypdf import PdfReader
    reader = PdfReader(path)
    pages, empty = [], 0
    for pg in reader.pages:
        t = pg.extract_text() or ""
        if not t.strip():
            empty += 1
        pages.append(t)
    return "\n".join(pages), empty, len(reader.pages)


def extract_xlsx(path, char_cap=5_000_000):
    """注意：不要用 read_only=True。部分 xlsx（WPS / 某些导出工具生成）
    在 sheet XML 里的 <dimension> 声明是错的（例如写成 A1），
    read_only 模式会信任它，导致整张表只读到 1 个单元格。
    这里用普通模式 + 字符上限，兼顾正确性与内存。"""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    parts, total = [], 0
    try:
        for ws in wb.worksheets:
            parts.append(f"### 工作表: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                cells = ["" if c is None else str(c) for c in row]
                if any(c.strip() for c in cells):
                    line = "\t".join(cells)
                    parts.append(line)
                    total += len(line)
                    if total >= char_cap:
                        parts.append("（已达字符上限，后续内容省略）")
                        return "\n".join(parts)
    finally:
        wb.close()
    return "\n".join(parts)


def extract(path):
    """抽取单个文件。返回 (文本 或 None, 类型标签, 备注)
    对压缩包：返回包内所有成员文本拼接后的结果。"""
    if is_archive(path):
        parts = []
        for e in iter_entries(path):
            if e["text"]:
                parts.append(e["text"])
        return "\n".join(parts), "压缩包", ""

    ext = os.path.splitext(path)[1].lower()

    if ext in (".docx", ".docm"):
        try:
            return extract_docx(path), "Word", ""
        except ImportError:
            return None, "Word", "缺少 python-docx，请 pip install python-docx"
        except Exception as e:
            return None, "Word", f"解析失败: {e}"

    if ext == ".pdf":
        try:
            text, empty, total = extract_pdf(path)
            note = f"{total} 页"
            if empty:
                note += f"，其中 {empty} 页无可提取文字（可能是扫描件/图片）"
            return text, "PDF", note
        except ImportError:
            return None, "PDF", "缺少 pypdf，请 pip install pypdf"
        except Exception as e:
            return None, "PDF", f"解析失败: {e}"

    if ext in (".xlsx", ".xlsm"):
        try:
            return extract_xlsx(path), "Excel", ""
        except ImportError:
            return None, "Excel", "缺少 openpyxl，请 pip install openpyxl"
        except Exception as e:
            return None, "Excel", f"解析失败: {e}"

    if ext == ".xls":
        return None, "Excel", "旧版 .xls 不支持，请另存为 .xlsx"
    if ext == ".doc":
        return None, "Word", "旧版 .doc 不支持，请另存为 .docx"

    if FORCE_TEXT:
        try:
            with open(path, "rb") as fh:
                return fh.read().decode("utf-8", errors="replace"), \
                    f"文本({ext})" if ext else "文本", "强制按文本读取"
        except Exception as e:
            return None, "未知", f"读取失败: {e}"

    try:
        text, enc = read_text_file(path)
        label = _label_for(ext)
        return text, label, "" if enc == "utf-8" else f"编码 {enc}"
    except BinaryFile as e:
        return None, "二进制", f"跳过（{e}）"
    except Exception as e:
        return None, "未知", f"无法按文本读取: {e}"


def _label_for(ext):
    if ext in KEIL_EXT:
        return "Keil"
    if ext in CODE_EXT:
        return "代码"
    if ext in TEXT_EXT or not ext:
        return "文本"
    return f"文本({ext})"


# ---------------------------------------------------------------- 压缩包遍历
def _iter_member(name, data, depth, max_depth):
    """name 是虚拟路径；嵌套压缩包用 ! 连接。
    reason 字段用于区分「可忽略的跳过」和「需要用户处理的提示」：
      artifact / skipdir / binary  → 只是跳过，汇总成数量即可
      error / need_tool / too_big  → 要显式告诉用户"""
    if depth < max_depth and is_archive(name):
        try:
            yield from _iter_archive_bytes(name, data, depth + 1, max_depth)
            return
        except Exception as e:
            yield {"name": name, "type": "压缩包", "text": None,
                   "note": f"嵌套压缩包无法打开: {e}", "reason": "error"}
            return

    # 编译产物：先按扩展名/目录/文件名判定，不解码，省时间
    if not INCLUDE_ARTIFACTS:
        if in_excluded(name):
            yield {"name": name, "type": "排除", "text": None,
                   "note": "跳过（--exclude 命中）", "reason": "excluded"}
            return
        if in_skip_dir(name):
            yield {"name": name, "type": "产物", "text": None,
                   "note": "跳过（产物目录）", "reason": "skipdir"}
            return
        if is_artifact(name):
            yield {"name": name, "type": "产物", "text": None,
                   "note": "跳过（编译产物）", "reason": "artifact"}
            return

    if FORCE_TEXT:
        try:
            yield {"name": name, "type": "文本", "text": data.decode("utf-8", errors="replace"),
                   "note": "强制按文本读取"}
        except Exception as e:
            yield {"name": name, "type": "未知", "text": None,
                   "note": f"读取失败: {e}", "reason": "error"}
        return

    try:
        text, enc = decode_bytes(base_name(name), data)
    except BinaryFile as e:
        yield {"name": name, "type": "二进制", "text": None,
               "note": f"跳过（{e}）", "reason": "binary"}
        return
    except Exception as e:
        yield {"name": name, "type": "未知", "text": None,
               "note": f"读取失败: {e}", "reason": "error"}
        return

    ext = os.path.splitext(base_name(name))[1].lower()
    yield {"name": name, "type": _label_for(ext), "text": text,
           "note": "" if enc == "utf-8" else f"编码 {enc}"}


def _iter_7z_bytes(virtual, data, depth, max_depth):
    """7z 用 py7zr 在内存里解，不落盘"""
    import py7zr
    from py7zr.io import BytesIOFactory

    with py7zr.SevenZipFile(io.BytesIO(data), "r") as z:
        dirs = {f.filename for f in z.list() if f.is_directory}
        factory = BytesIOFactory(limit=512 * 1024 * 1024)
        z.extractall(factory=factory)
        for name, bio in factory.products.items():
            if name in dirs:
                continue
            yield from _iter_member(f"{virtual}!{name}", bio.read(), depth, max_depth)


def _iter_rar_path(path, depth, max_depth):
    """rar 只能借外部工具，先解到临时目录再遍历"""
    size = os.path.getsize(path)
    base = os.path.basename(path)
    if size > RAR_EXTRACT_LIMIT:
        yield {"name": base, "type": "压缩包", "text": None,
               "note": f"文件 {size / 1048576:.0f} MB 过大，请手动解压后再估算",
               "reason": "too_big"}
        return

    dest, err = extract_rar_to_temp(path)
    if dest is None:
        yield {"name": base, "type": "压缩包", "text": None, "note": err,
               "reason": "need_tool"}
        return
    try:
        for root, _, names in os.walk(dest):
            for n in sorted(names):
                full = os.path.join(root, n)
                rel = os.path.relpath(full, dest).replace("\\", "/")
                try:
                    with open(full, "rb") as fh:
                        data = fh.read()
                except OSError as e:
                    yield {"name": rel, "type": "未知", "text": None,
                           "note": f"读取失败: {e}", "reason": "error"}
                    continue
                yield from _iter_member(rel, data, depth, max_depth)
    finally:
        shutil.rmtree(dest, ignore_errors=True)


def _iter_archive_bytes(virtual, data, depth, max_depth):
    """从内存里的字节打开嵌套压缩包"""
    kind = archive_kind(virtual)

    if kind == "zip":
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            for info in zf.infolist():
                if info.is_dir():
                    continue
                yield from _iter_member(f"{virtual}!{info.filename}", zf.read(info),
                                        depth, max_depth)
        return

    if kind == "7z":
        yield from _iter_7z_bytes(virtual, data, depth, max_depth)
        return

    if kind == "rar":
        # rar 没有纯 Python 解压，只能先落到临时文件再借外部工具
        fd, tmp = tempfile.mkstemp(suffix=".rar", prefix="toknested_")
        try:
            with os.fdopen(fd, "wb") as fh:
                fh.write(data)
            yield from _iter_rar_path(tmp, depth, max_depth)
        finally:
            try:
                os.remove(tmp)
            except OSError:
                pass
        return

    with tarfile.open(fileobj=io.BytesIO(data)) as tf:
        for m in tf.getmembers():
            if not m.isfile():
                continue
            fh = tf.extractfile(m)
            yield from _iter_member(f"{virtual}!{m.name}", fh.read() if fh else b"",
                                    depth, max_depth)


def iter_entries(path, max_depth=3):
    """统一遍历入口：产出 {name, type, text, note}
    - 普通文件 → 一条
    - 压缩包   → 包内每个成员一条（可递归嵌套）"""
    kind = archive_kind(path)
    if kind is None:
        text, label, note = extract(path)
        yield {"name": os.path.basename(path), "type": label, "text": text, "note": note}
        return

    if kind == "rar":
        yield from _iter_rar_path(path, 0, max_depth)
        return

    try:
        if kind == "zip":
            with zipfile.ZipFile(path) as zf:
                for info in zf.infolist():
                    if info.is_dir():
                        continue
                    yield from _iter_member(info.filename, zf.read(info), 0, max_depth)
        elif kind == "7z":
            with open(path, "rb") as fh:
                yield from _iter_7z_bytes(os.path.basename(path), fh.read(), 0, max_depth)
        else:  # tar 系
            with tarfile.open(path) as tf:
                for m in tf.getmembers():
                    if not m.isfile():
                        continue
                    fh = tf.extractfile(m)
                    yield from _iter_member(m.name, fh.read() if fh else b"", 0, max_depth)
    except Exception as e:
        yield {"name": os.path.basename(path), "type": "压缩包",
               "text": None, "note": f"无法打开压缩包: {e}"}


# ---------------------------------------------------------------- 主流程
def collect_files(paths):
    """返回 [(路径, 是否用户明确指定)]。
    目录扫描出来的文件 explicit=False；命令行直接点名的 explicit=True。
    区别在于：目录扫描时编译产物默认跳过，用户点名的文件则照算。"""
    out, dup_warn = [], []
    for p in paths:
        if os.path.isdir(p):
            for root, dirs, names in os.walk(p):
                dirset = {d.lower() for d in dirs}
                for n in sorted(names):
                    full = os.path.join(root, n)
                    out.append((full, False))
                    # 同级同时存在 X.zip 和 X/ 目录 → 大概率是同一份内容的两份拷贝
                    if is_archive(n):
                        stem = n.lower()
                        for e in sorted(ARCHIVE_EXT, key=len, reverse=True):
                            if stem.endswith(e):
                                stem = stem[: -len(e)]
                                break
                        if stem and stem in dirset:
                            dup_warn.append(os.path.join(root, n))
        elif os.path.isfile(p):
            out.append((p, True))
        else:
            print(f"⚠️  找不到：{p}", file=sys.stderr)

    for w in dup_warn:
        print(f"⚠️  {w}", file=sys.stderr)
        print("    同级有同名文件夹，内容可能是同一份，会重复计算。", file=sys.stderr)
        print("    建议只指定压缩包，或只指定文件夹。", file=sys.stderr)
    return out


def _dir_rollup(detail, depth=2, top_n=12):
    """按目录汇总 token，帮用户看清「大头在哪」。
    detail: [(绝对路径, 字符数, token)]，返回 [(显示名, 文件数, 字符数, token)]"""
    if not detail:
        return []
    try:
        root = os.path.commonpath([os.path.dirname(p) for p, _, _ in detail])
    except ValueError:
        root = ""
    groups = {}
    for path, ch, tk in detail:
        rel = os.path.relpath(path, root) if root else path
        parts = rel.replace("\\", "/").split("/")[:-1]
        key = "/".join(parts[:depth]) if parts else "(根目录)"
        g = groups.setdefault(key, {"files": 0, "char": 0, "tok": 0})
        g["files"] += 1
        g["char"] += ch
        g["tok"] += tk
    rows = sorted(groups.items(), key=lambda kv: -kv[1]["tok"])
    return [(k, v["files"], v["char"], v["tok"]) for k, v in rows[:top_n]]


def analyze(paths, tok, all_files=False, top_n=10):
    """返回 (报告 dict)。报告含：分组行、普通文件行、总 token 等。
    注意：是否计入编译产物由模块级 INCLUDE_ARTIFACTS 控制（对应 --all）。"""
    archives, plain = [], []
    for path, explicit in collect_files(paths):
        if is_archive(path):
            archives.append(path)
        else:
            plain.append((path, explicit))

    # --- 普通文件 ---
    plain_rows, plain_tok, plain_char, plain_skipped = [], 0, 0, 0
    plain_detail, excluded_n = [], 0
    for path, explicit in plain:
        name = os.path.basename(path)

        # --exclude 命中的目录，跳过
        if in_excluded(path):
            excluded_n += 1
            plain_skipped += 1
            plain_rows.append((name, "排除", "-", "-", "跳过（--exclude 命中）", None))
            continue

        # 目录扫描出来的编译产物默认跳过；用户点名的文件则照算
        if not INCLUDE_ARTIFACTS and not explicit and (in_skip_dir(path) or is_artifact(path)):
            plain_skipped += 1
            plain_rows.append((name, "产物", "-", "-", "跳过（编译产物）", None))
            continue

        size = os.path.getsize(path)
        text, label, note = extract(path)
        if text is None:
            plain_skipped += 1
            plain_rows.append((name, label, "-", "-", note, None))
            continue
        n = count(tok, text)
        plain_tok += n
        plain_char += len(text)
        if size > 0 and not text:
            note = (note + "；" if note else "") + "未提取到文字"
        plain_rows.append((name, label, human(len(text)), human(n), note, n))
        plain_detail.append((path, len(text), n))

    # --- 压缩包 ---
    arch_reports = []
    for path in archives:
        members = list(iter_entries(path))
        included, warnings = [], []
        skipped_art = skipped_bin = skipped_other = skipped_exc = 0
        for e in members:
            if e["text"] is None:
                r = e.get("reason")
                if r in ("artifact", "skipdir"):
                    skipped_art += 1
                elif r == "excluded":
                    skipped_exc += 1
                elif r == "binary":
                    skipped_bin += 1
                else:
                    skipped_other += 1
                    if e.get("note"):
                        warnings.append((e["name"], e["note"]))
                continue
            if not e["text"].strip():
                continue
            included.append({**e, "tok": count(tok, e["text"]), "char": len(e["text"])})

        # 按扩展名汇总
        groups = {}
        for e in included:
            ext = os.path.splitext(base_name(e["name"]))[1].lower() or "(无扩展名)"
            g = groups.setdefault(ext, {"files": 0, "char": 0, "tok": 0})
            g["files"] += 1
            g["char"] += e["char"]
            g["tok"] += e["tok"]
        total_tok = sum(g["tok"] for g in groups.values())
        total_char = sum(g["char"] for g in groups.values())

        top = sorted(included, key=lambda e: -e["tok"])[:top_n]
        # 整个压缩包都没能打开（例如缺 unrar/7-Zip）时，别把那条错误记录算成「包内条目」
        opened = not any(e.get("reason") == "need_tool" for e in members)
        arch_reports.append({
            "path": path, "name": os.path.basename(path), "opened": opened,
            "members": len(members) if opened else 0, "included": len(included),
            "skipped_art": skipped_art, "skipped_bin": skipped_bin,
            "skipped_other": skipped_other, "skipped_exc": skipped_exc,
            "groups": sorted(groups.items(), key=lambda kv: -kv[1]["tok"]),
            "top": top, "files": included if all_files else None,
            "warnings": warnings,
            "total_tok": total_tok, "total_char": total_char,
        })

    return {
        "plain_rows": plain_rows, "plain_tok": plain_tok, "plain_char": plain_char,
        "plain_skipped": plain_skipped, "plain_count": len(plain),
        "plain_included": len(plain) - plain_skipped,
        "plain_detail": plain_detail, "dir_rows": _dir_rollup(plain_detail),
        "excluded_n": excluded_n,
        "archives": arch_reports,
        "total_tok": plain_tok + sum(a["total_tok"] for a in arch_reports),
        "total_char": plain_char + sum(a["total_char"] for a in arch_reports),
        "file_count": len(plain) + len(archives),
    }


def main():
    ap = argparse.ArgumentParser(description="上传文件前估算 token 用量")
    ap.add_argument("paths", nargs="*", help="一个或多个文件 / 目录（支持 .zip 等压缩包）")
    ap.add_argument("--window", type=int, default=128000,
                    help="参考上下文窗口大小，默认 128000（仅用于算占比）")
    ap.add_argument("-q", "--quiet", action="store_true", help="只输出总 token 数")
    ap.add_argument("--detail", action="store_true", help="附上每个文件的开头预览")
    ap.add_argument("--all", action="store_true",
                    help="连编译产物（.o/.axf/.map/Objects 目录等）一起算")
    ap.add_argument("--all-files", action="store_true",
                    help="逐个文件列出（压缩包内逐个成员、目录下逐个文件）")
    ap.add_argument("--exclude", action="append", default=[], metavar="目录名",
                    help="跳过路径中名为该目录名的内容，可重复。"
                         "例：--exclude Driver --exclude CMSIS")
    ap.add_argument("--top", type=int, default=10, help="汇总时列出最大的 N 个文件，默认 10")
    ap.add_argument("--force-text", action="store_true",
                    help="忽略二进制嗅探，把文件一律当文本读")
    args = ap.parse_args()

    global FORCE_TEXT, INCLUDE_ARTIFACTS, EXCLUDE_PATS
    FORCE_TEXT = args.force_text
    INCLUDE_ARTIFACTS = args.all
    EXCLUDE_PATS = [p.strip().lower() for p in args.exclude if p.strip()]

    if not args.paths:
        ap.print_help()
        print("\n提示：也可以把文件直接拖到 drag_files_here.cmd 上。")
        return

    tok, backend = load_tokenizer()
    rep = analyze(args.paths, tok, all_files=args.all_files, top_n=args.top)

    if args.quiet:
        print(rep["total_tok"])
        return

    print(f"后端: {backend}")
    print()

    # ---- 普通文件 ----
    if rep["plain_rows"]:
        if len(rep["plain_rows"]) > 30 and not args.all_files:
            # 文件太多时，先给「哪个目录吃掉了 token」，避免刷屏
            W, A = (44, 8, 12, 12, 8), ("left", "right", "right", "right", "right")
            rows = []
            for d, files, ch, tk in rep["dir_rows"]:
                pct = tk / rep["plain_tok"] * 100 if rep["plain_tok"] else 0
                rows.append((d, files, human(ch), human(tk), f"{pct:.1f}%"))
            table(("目录", "文件数", "字符数", "Token", "占比"), W, A, rows)
            print(f"  （共 {len(rep['plain_rows'])} 个文件，按目录汇总；"
                  f"加 --all-files 看逐个文件）")
            print()
        else:
            W, A = (32, 9, 10, 10, 34), ("left", "left", "right", "right", "left")
            table(("文件", "类型", "字符数", "Token", "备注"), W, A,
                  [(n, l, c, t, note) for n, l, c, t, note, _ in rep["plain_rows"]])
            print()

    # ---- 压缩包 ----
    for a in rep["archives"]:
        print(f"压缩包: {a['name']}")
        if not a["opened"]:
            for name, note in a["warnings"][:5]:
                print(f"  ⚠️  {name}: {note}")
            print()
            continue
        info = f"  包内 {a['members']} 个条目，计入 {a['included']} 个"
        if a["skipped_art"]:
            info += f"，跳过编译产物 {a['skipped_art']} 个"
        if a["skipped_bin"]:
            info += f"，跳过二进制 {a['skipped_bin']} 个"
        if a.get("skipped_exc"):
            info += f"，按 --exclude 跳过 {a['skipped_exc']} 个"
        if a["skipped_other"]:
            info += f"，其他 {a['skipped_other']} 个"
        print(info)
        print()

        if not a["included"]:
            print("  （没有可计入的文本内容）")
            for name, note in a["warnings"][:5]:
                print(f"  ⚠️  {name}: {note}")
            print()
            continue

        if a["files"] is not None:
            # --all-files：逐个列出
            W, A = (56, 10, 12), ("left", "right", "right")
            table(("包内文件", "字符数", "Token"), W, A,
                  [(e["name"], human(e["char"]), human(e["tok"])) for e in a["files"]],
                  indent=2)
        else:
            # 默认：按扩展名汇总
            W, A = (18, 8, 12, 12, 8), ("left", "right", "right", "right", "right")
            rows = []
            for ext, g in a["groups"]:
                pct = g["tok"] / a["total_tok"] * 100 if a["total_tok"] else 0
                rows.append((ext, g["files"], human(g["char"]), human(g["tok"]), f"{pct:.1f}%"))
            table(("扩展名", "文件数", "字符数", "Token", "占比"), W, A, rows, indent=2)
            print()

            # 最大的 N 个文件
            if a["top"]:
                print(f"  最大的 {len(a['top'])} 个文件：")
                W2, A2 = (54, 12, 8), ("left", "right", "right")
                trows = []
                for e in a["top"]:
                    pct = e["tok"] / a["total_tok"] * 100 if a["total_tok"] else 0
                    trows.append((e["name"], human(e["tok"]), f"{pct:.1f}%"))
                table(("包内文件", "Token", "占比"), W2, A2, trows, indent=2)

        if a["warnings"]:
            print()
            for name, note in a["warnings"][:5]:
                print(f"  ⚠️  {name}")
                print(f"      {note}")
            if len(a["warnings"]) > 5:
                print(f"  ……还有 {len(a['warnings']) - 5} 条类似提示")

        print()
        print(f"  小计：字符 {human(a['total_char'])}，Token {human(a['total_tok'])}")
        print()

    # ---- 总计 ----
    print("=" * 66)
    parts = []
    if rep["plain_count"]:
        s = f"普通文件 {rep['plain_count']} 个"
        if rep["plain_skipped"]:
            s += f"（计入 {rep['plain_included']}）"
        parts.append(s)
    if rep["archives"]:
        n = sum(a["included"] for a in rep["archives"])
        nfail = sum(1 for a in rep["archives"] if not a["opened"])
        s = f"压缩包 {len(rep['archives'])} 个（包内计入 {n}）"
        if nfail:
            s += f"，其中 {nfail} 个打不开"
        parts.append(s)
    line = "，".join(parts) if parts else "无输入"
    print(f"合计：{line}　字符 {human(rep['total_char'])}　Token {human(rep['total_tok'])}")

    if rep["total_tok"]:
        pct = rep["total_tok"] / args.window * 100
        msg = f"占参考窗口 {human(args.window)} 的 {pct:.1f}%"
        if pct > 100:
            msg += "   ⚠️  超出窗口，需要拆分或精简"
        elif pct > 60:
            msg += "   ⚠️  占用偏高，可能挤压对话空间"
        print(msg)
    print("=" * 66)

    if args.detail:
        print()
        for path, explicit in collect_files(args.paths):
            if is_archive(path):
                continue
            if not INCLUDE_ARTIFACTS and not explicit and (
                    in_excluded(path) or in_skip_dir(path) or is_artifact(path)):
                continue
            text, label, _ = extract(path)
            if not text:
                continue
            head = text[:200].replace("\n", " ⏎ ")
            print(f"── {os.path.basename(path)} ──")
            print(f"   {head}{'…' if len(text) > 200 else ''}")
            print()


if __name__ == "__main__":
    main()
